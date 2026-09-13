"""
scripts/generate_pptx_deck.py -- builds docs/Predicting_Drive_Failure_Slides.pptx,
a defense-ready slide deck.

Pulls real numbers directly from model/metrics.json, model/survival_metrics.json,
and data/provenance.json, same as generate_docx_report.py -- re-run this after
any retrain to regenerate the deck with fresh numbers.

Usage:
    python scripts/generate_pptx_deck.py
"""

import json
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ACCENT = RGBColor(0x0D, 0x7D, 0x6F)
INK = RGBColor(0x16, 0x21, 0x1F)
MUTED = RGBColor(0x55, 0x5F, 0x5C)
GOOD = RGBColor(0x2F, 0x8A, 0x5C)
WARN = RGBColor(0xA5, 0x6A, 0x12)
CRIT = RGBColor(0xB8, 0x45, 0x2C)
PAPER = RGBColor(0xF7, 0xF9, 0xF8)

OUT_PATH = "docs/Predicting_Drive_Failure_Slides.pptx"
SCREENSHOTS = "docs/screenshots"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def load_json(path):
    with open(path) as f:
        return json.load(f)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    return slide


def add_title(slide, text, size=32, color=None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color or ACCENT
    return box


def add_kicker(slide, text):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.05), Inches(10), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text.upper()
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = MUTED
    return box


def add_bullets(slide, items, left=0.6, top=1.4, width=12.1, height=5.5, size=18, color=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, sub = item
        else:
            text, sub = item, None
        run = p.add_run()
        run.text = f"▪  {text}"
        run.font.size = Pt(size)
        run.font.color.rgb = color or INK
        p.space_after = Pt(10)
        if sub:
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = f"     {sub}"
            r2.font.size = Pt(size - 4)
            r2.font.italic = True
            r2.font.color.rgb = MUTED
            p2.space_after = Pt(14)
    return box


def add_footer(slide, text="Predicting Drive Failure — IT Capstone"):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.1), Inches(10), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED


def add_stat_row(slide, stats, top=3.2, box_w=2.8, box_h=1.5, gap=0.25):
    n = len(stats)
    total_w = n * box_w + (n - 1) * gap
    left = (13.333 - total_w) / 2
    for label, value, sub in stats:
        box = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(box_w), Inches(box_h))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        box.line.color.rgb = RGBColor(0xD7, 0xDD, 0xD7)
        box.line.width = Pt(1)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.1)
        p0 = tf.paragraphs[0]
        p0.alignment = PP_ALIGN.CENTER
        r0 = p0.add_run()
        r0.text = label.upper()
        r0.font.size = Pt(11)
        r0.font.color.rgb = MUTED
        p1 = tf.add_paragraph()
        p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run()
        r1.text = str(value)
        r1.font.size = Pt(26)
        r1.font.bold = True
        r1.font.color.rgb = INK
        if sub:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = sub
            r2.font.size = Pt(9)
            r2.font.color.rgb = MUTED
        left += box_w + gap


def add_picture_slide(prs, kicker, title, image, caption):
    slide = blank_slide(prs)
    add_kicker(slide, kicker)
    add_title(slide, title)
    img_path = os.path.join(SCREENSHOTS, image)
    if os.path.exists(img_path):
        pic = slide.shapes.add_picture(img_path, Inches(1.2), Inches(1.5), width=Inches(10.9))
        # cap height so it doesn't run off the slide; shrink proportionally if needed
        if pic.height > Emu(int(Inches(5.3))):
            ratio = Inches(5.3) / pic.height
            pic.height = Inches(5.3)
            pic.width = int(pic.width * ratio)
            pic.left = int((SLIDE_W - pic.width) / 2)
    cap = slide.shapes.add_textbox(Inches(0.6), Inches(6.85), Inches(12.1), Inches(0.4))
    p = cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = caption
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = MUTED
    add_footer(slide)
    return slide


def main():
    metrics = load_json("model/metrics.json")
    surv = load_json("model/survival_metrics.json")
    prov = load_json("data/provenance.json")

    report = metrics["classification_report"]["1"]
    cv = metrics.get("cross_fold_stability", {})
    cv_mean = cv.get("mean", {})
    cv_std = cv.get("std", {})
    model_name = ", ".join(prov.get("models_filtered") or ["(all models)"])
    date_range = prov.get("date_range") or ["?", "?"]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ---------- 1. Title ----------
    slide = blank_slide(prs)
    box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.3), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Predicting Drive Failure"
    r.font.size = Pt(48)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "A Predictive-Maintenance System Built on Real Backblaze Drive Stats"
    r2.font.size = Pt(20)
    r2.font.color.rgb = MUTED
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = "\nIT Capstone Project"
    r3.font.size = Pt(14)
    r3.font.color.rgb = INK

    # ---------- 2. Problem & Motivation ----------
    slide = blank_slide(prs)
    add_kicker(slide, "The Problem")
    add_title(slide, "Unplanned Drive Failure Is an Operational Cost")
    add_bullets(slide, [
        "Reactive maintenance (replace only after failure) risks data loss and unplanned downtime",
        "Blanket proactive replacement wastes budget on drives that would have run fine",
        "The real industrial answer: predict early, estimate remaining useful life, and pick a threshold suited to your own cost structure",
    ])
    add_footer(slide)

    # ---------- 3. Why Different ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Positioning")
    add_title(slide, "Not Just Another \"Predict Drive Failure\" Notebook")
    add_bullets(slide, [
        ("Early-warning labeling, not same-day diagnosis", "Same-day features already look like a failure — trivial and non-predictive; excluded from training"),
        ("Remaining useful life, not just yes/no", "A Cox Proportional Hazards survival model, the standard real-world framing"),
        ("Per-prediction explainability", "SHAP explains THIS drive's score, not just global feature importance"),
        ("A genuine decision-support result", "The right cost threshold depends on your own cost structure — no universal answer"),
    ], size=17)
    add_footer(slide)

    # ---------- 4. Data ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Data")
    add_title(slide, "Real Backblaze Drive Stats, One Quarter")
    add_stat_row(slide, [
        ("Drive Model", model_name, "highest-volume model with full SMART coverage"),
        ("Quarter", f"{date_range[0][:7]}", f"{date_range[0]} to {date_range[1]}"),
        ("Real Failures", prov.get("n_failure_rows", "?"), "early-warning-positive rows"),
    ], top=1.7, box_w=3.6)
    add_bullets(slide, [
        "A real vendor-reporting discovery during model selection: the highest-volume drive model reported ZERO values for 2 of the 7 SMART attributes this project relies on — a genuine WD/Toshiba-vs-Seagate reporting difference, confirmed while picking a model",
    ], top=3.6, size=16)
    add_footer(slide)

    # ---------- 5. Classifier ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Methodology")
    add_title(slide, "Classifier: RandomForest, Grouped Split")
    add_bullets(slide, [
        ("RandomForestClassifier", "300 trees, max depth 8, class_weight='balanced', on 7 SMART features"),
        ("Split grouped by drive, not by row", "StratifiedGroupKFold on serial_number — no drive's rows appear on both sides of train/test"),
        ("Why this matters", "A failing drive contributes several rows to the positive class; a naive row-level split let the same physical drive leak across train and test, inflating early metrics — caught and fixed"),
    ], size=17)
    add_footer(slide)

    # ---------- 6. Results: Classifier ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Results")
    add_title(slide, "Classifier Performance (Real Held-Out Test Fold)")
    add_stat_row(slide, [
        ("Precision", f"{report['precision']:.0%}", "failure class"),
        ("Recall", f"{report['recall']:.0%}", "failure class"),
        ("F1", f"{report['f1-score']:.0%}", "harmonic mean"),
        ("ROC-AUC", f"{metrics.get('roc_auc', 0):.3f}", "test set"),
    ], top=1.9, box_w=2.8)
    add_bullets(slide, [
        (f"{metrics.get('n_drives_train','?')} train / {metrics.get('n_drives_test','?')} test drives",
         "zero drive overlap between the two sets"),
    ], top=3.8, size=16)
    add_footer(slide)

    # ---------- 7. Cross-fold stability ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Statistical Honesty")
    add_title(slide, "Report the Range, Not Just One Split")
    add_bullets(slide, [
        ("Only ~43 real failing drives underlie every classifier metric",
         "genuinely few for a held-out evaluation — which fold you land in matters a lot"),
        (f"Precision stable across folds: {cv_mean.get('precision',0):.2f} ± {cv_std.get('precision',0):.2f}", None),
        (f"Recall unstable across folds: {cv_mean.get('recall',0):.2f} ± {cv_std.get('recall',0):.2f} (range ~0.62–1.00)",
         "depends heavily on which specific failing drives land in the test fold"),
        ("All 4 folds evaluated and disclosed in model/metrics.json, not just the deployed one", None),
    ], size=17)
    add_footer(slide)

    # ---------- 8. Survival model ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Methodology")
    add_title(slide, "Survival Analysis: Remaining Useful Life")
    add_stat_row(slide, [
        ("Concordance", f"{surv['concordance_index']:.3f}", "0.5=random, 1.0=perfect"),
        ("Drives Modeled", f"{surv['n_drives']:,}", ""),
        ("Failure Events", surv["n_events"], "genuinely few — disclosed"),
    ], top=1.9, box_w=3.4)
    add_bullets(slide, [
        ("Cox Proportional Hazards model (lifelines, penalizer=0.1)",
         "estimates days remaining per drive, not just a risk label"),
        ("Kaplan-Meier baseline", "fleet-wide reference survival curve shown alongside each drive's own curve"),
    ], top=3.8, size=16)
    add_footer(slide)

    # ---------- 9. Dashboards intro ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Product")
    add_title(slide, "Three Dashboards, Three Audiences")
    add_bullets(slide, [
        ("Fleet Overview — leadership/budget", "3D rack, event feed, live-feel replay, cost-avoidance estimate"),
        ("Operator Lookup — technician", "per-drive SHAP explanation, not just a global importance chart"),
        ("Optimization Lab — decision-maker", "remaining useful life + 3D cost-optimized threshold surface"),
    ], size=19)
    add_footer(slide)

    # ---------- 10-12. Screenshots ----------
    add_picture_slide(prs, "Dashboard", "Fleet Overview", "03_fleet_overview_rack.png",
                       "3D rack colored by risk tier, real trend line, event feed with both elevated and critical events")
    add_picture_slide(prs, "Dashboard", "Operator Lookup", "04_operator_lookup.png",
                       "A real high-risk drive (79%) with its SHAP explanation — which readings actually drove this score")
    add_picture_slide(prs, "Dashboard", "Optimization Lab", "06_optimization_lab_cost_surface.png",
                       "3D cost surface: the optimal alert threshold moves with the organization's cost ratio")

    # ---------- 13. Automation ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Automation")
    add_title(slide, "Two-Tier Alerting, One Layer Genuinely Unattended")
    add_bullets(slide, [
        ("⚠️ Early warning", "fires the moment a drive crosses into ELEVATED risk — real lead time"),
        ("🚨 Critical", "fires when a drive crosses into CRITICAL risk — act now"),
        ("In-browser alert", "instant, while someone is actively using Fleet Overview"),
        ("Scheduled GitHub Actions check", "runs hourly, independent of the app or anyone watching — free and unlimited on a public repo, the genuinely unattended piece"),
    ], size=17)
    add_footer(slide)

    # ---------- 14. Limitations ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Honesty")
    add_title(slide, "Known Limitations")
    add_bullets(slide, [
        "Rare-event target (~1.36% real annual failure rate) — report precision/recall/F1, not accuracy",
        "Only ~43 real failing drives — cross-fold range reported, not a single flattering split",
        "One vendor's fleet — a demonstrated method, not a drop-in production tool",
        "“Live” is a labeled replay of real historical data, not a live sensor feed",
        "Cost-avoidance figures are adjustable planning assumptions, not sourced statistics",
        "Static, last-observed covariates in the survival model, not full time-varying history",
    ], size=15)
    add_footer(slide)

    # ---------- 15. Future work ----------
    slide = blank_slide(prs)
    add_kicker(slide, "Roadmap")
    add_title(slide, "Future Work")
    add_bullets(slide, [
        "Validate the Cox proportional-hazards assumption",
        "Calibration curve / Brier score for predicted probabilities",
        "Extend GitHub Actions automation to periodic model retraining",
        "Data drift detection between training and live data",
        "A pytest test suite and CI workflow",
    ], size=18)
    add_footer(slide)

    # ---------- 16. Thank you ----------
    slide = blank_slide(prs)
    box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11.3), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Thank You"
    r.font.size = Pt(44)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Questions?"
    r2.font.size = Pt(20)
    r2.font.color.rgb = MUTED
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = "\nhttps://hdd-failure-predictor.onrender.com\nhttps://github.com/dp2426-NAU/hdd-failure-predictor"
    r3.font.size = Pt(13)
    r3.font.color.rgb = INK

    os.makedirs("docs", exist_ok=True)
    prs.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
